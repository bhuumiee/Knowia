
import json
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

BASE_DIR = Path(__file__).parent
UPLOAD_DIR = BASE_DIR / "uploads"
CACHE_PATH = BASE_DIR / "material_cache.json"

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}


def extract_text(file_path):
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        reader = PdfReader(str(file_path))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if extension == ".docx":
        document = Document(str(file_path))
        return "\n".join(
            paragraph.text for paragraph in document.paragraphs
        )

    if extension == ".pptx":
        presentation = Presentation(str(file_path))
        slide_texts = []

        for number, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text.strip()
            ]
            if texts:
                slide_texts.append(
                    f"Slide {number}:\n" + "\n".join(texts)
                )

        return "\n\n".join(slide_texts)

    if extension == ".txt":
        return file_path.read_text(encoding="utf-8", errors="replace")

    return ""


def save_cache(cache):
    temporary_path = CACHE_PATH.with_suffix(".tmp")
    temporary_path.write_text(
        json.dumps(cache, indent=2, ensure_ascii=False),
        encoding="utf-8"
    )
    temporary_path.replace(CACHE_PATH)


if CACHE_PATH.exists():
    try:
        cache = json.loads(CACHE_PATH.read_text(encoding="utf-8"))
        if not isinstance(cache, dict):
            cache = {}
    except (json.JSONDecodeError, OSError):
        cache = {}
else:
    cache = {}

files = sorted(
    path for path in UPLOAD_DIR.iterdir()
    if path.is_file() and not path.name.endswith(".part")
)

for index, file_path in enumerate(files, start=1):
    existing = cache.get(file_path.name, {})

    if (
        "text_extracted" in existing
        and "text_preview" in existing
    ):
        print(f"[{index}/{len(files)}] Cached: {file_path.name}")
        continue

    try:
        if file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            text = extract_text(file_path)
        else:
            text = ""

        cache[file_path.name] = {
            "text_extracted": bool(text.strip()),
            "text_preview": text[:300],
            "text_length": len(text)
        }

        save_cache(cache)
        print(
            f"[{index}/{len(files)}] Done: {file_path.name} "
            f"({len(text)} characters)"
        )

    except Exception as error:
        cache[file_path.name] = {
            "text_extracted": False,
            "text_preview": "",
            "text_length": 0
        }
        save_cache(cache)
        print(f"[{index}/{len(files)}] Failed: {file_path.name}: {error}")

print("Backfill finished.")
