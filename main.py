import os
from dotenv import load_dotenv
from pathlib import Path
from fastapi import FastAPI, UploadFile, File, HTTPException
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from fastapi.middleware.cors import CORSMiddleware
import re
from pydantic import BaseModel
from google import genai
from youtube_transcript_api import YouTubeTranscriptApi
from urllib.request import Request, urlopen
from html.parser import HTMLParser
from urllib.parse import urlparse, parse_qs
import time
import tempfile
import yt_dlp
from faster_whisper import WhisperModel
import json
import pymupdf
from PIL import Image
import io
# Learner model storage
LEARNER_STORE_PATH = Path(__file__).parent / "learner_store.json"


def load_learner_store():
    if not LEARNER_STORE_PATH.exists():
        return {}

    try:
        with open(LEARNER_STORE_PATH, "r", encoding="utf-8") as file:
            data = json.load(file)

        return data if isinstance(data, dict) else {}

    except (json.JSONDecodeError, OSError):
        return {}


def save_learner_store(store):
    with open(LEARNER_STORE_PATH, "w", encoding="utf-8") as file:
        json.dump(store, file, indent=2, ensure_ascii=False)

def get_or_create_learner(store, learner_id):
    if learner_id not in store:
        store[learner_id] = {
            "learnerId": learner_id,
            "isNewLearner": True,
            "topicMastery": {},
            "quizHistory": [],
            "conversationHistory": [],
            "createdAt": time.strftime("%Y-%m-%dT%H:%M:%S"),
            "updatedAt": time.strftime("%Y-%m-%dT%H:%M:%S")
        }

    return store[learner_id]


def update_topic_mastery(learner, topic, is_correct, evidence_type):
    topic = (topic or "General").strip() or "General"

    mastery = learner["topicMastery"].setdefault(
        topic,
        {
            "estimate": 50.0,
            "attempts": 0,
            "correct": 0,
            "incorrect": 0,
            "lastEvidence": None
        }
    )

    # A simple, explainable estimate that moves with each
    # observed correct or incorrect answer.
    old_estimate = mastery["estimate"]
    learning_rate = 0.15

    if is_correct:
        new_estimate = old_estimate + learning_rate * (100 - old_estimate)
        mastery["correct"] += 1
    else:
        new_estimate = old_estimate - learning_rate * old_estimate
        mastery["incorrect"] += 1

    mastery["estimate"] = round(
        max(0.0, min(100.0, new_estimate)), 2
    )
    mastery["attempts"] += 1
    mastery["lastEvidence"] = evidence_type

    learner["isNewLearner"] = False
    learner["updatedAt"] = time.strftime("%Y-%m-%dT%H:%M:%S")

    return mastery        

load_dotenv()

app = FastAPI(title="Knowvia API")

# Allow the React frontend to communicate with FastAPI
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://localhost:5174",],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Folder where uploaded files will be stored
UPLOAD_DIR = Path(__file__).parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
MATERIAL_CACHE_PATH = Path(__file__).parent / "material_cache.json"

def load_material_cache():
    if not MATERIAL_CACHE_PATH.exists():
        return {}

    try:
        with MATERIAL_CACHE_PATH.open("r", encoding="utf-8") as file:
            data = json.load(file)
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def save_material_cache(data):
    with MATERIAL_CACHE_PATH.open("w", encoding="utf-8") as file:
        json.dump(data, file, indent=2, ensure_ascii=False)


ALLOWED_EXTENSIONS = {
    ".pdf", ".ppt", ".pptx", ".doc", ".docx",
    ".png", ".jpg", ".jpeg", ".webp",
    ".mp4", ".mov", ".avi", ".webm", ".mkv"
}

MAX_FILE_SIZE = 200 * 1024 * 1024  # 200 MB

def get_youtube_video_id(url: str):
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower()
    
    if host in {"youtu.be", "www.youtu.be"}:
        video_id = parsed.path.strip("/").split("/")[0]
    elif host in {"youtube.com", "www.youtube.com", "m.youtube.com"}:
        if parsed.path == "/watch":
            video_id = parse_qs(parsed.query).get("v", [""])[0]
        elif parsed.path.startswith(("/shorts/", "/embed/")):
            parts = parsed.path.strip("/").split("/")
            video_id = parts[1] if len(parts) > 1 else ""
        else:
            return None
    else:
        return None

    if video_id and all(
        char.isalnum() or char in "_-" for char in video_id
    ):
        return video_id

    return None


def transcribe_youtube_audio(video_id: str) -> str:
    """Download YouTube audio and transcribe it with Whisper."""

    ffmpeg_dir = (
        r"C:\Users\admin\Downloads\ffmpeg-9.0.2-essentials_build"
        r"\ffmpeg-9.0.2-essentials_build\bin"
    )

    video_url = f"https://www.youtube.com/watch?v={video_id}"

    with tempfile.TemporaryDirectory() as temp_dir:
        output_template = str(Path(temp_dir) / "audio.%(ext)s")

        ydl_options = {
            "format": "bestaudio/best",
            "outtmpl": output_template,
            "noplaylist": True,
            "quiet": True,
            "no_warnings": True,
            "ffmpeg_location": ffmpeg_dir,
            "postprocessors": [
                {
                    "key": "FFmpegExtractAudio",
                    "preferredcodec": "mp3",
                    "preferredquality": "128",
                }
            ],
        }

        with yt_dlp.YoutubeDL(ydl_options) as downloader:
            downloader.download([video_url])

        audio_path = Path(temp_dir) / "audio.mp3"

        if not audio_path.is_file():
            raise ValueError("Could not extract audio from the video.")

        model = WhisperModel(
            "base",
            device="cpu",
            compute_type="int8",
        )

        segments, _ = model.transcribe(
            str(audio_path),
            vad_filter=True,
        )

        transcript = "\n".join(
            segment.text.strip()
            for segment in segments
            if segment.text.strip()
        )

        if not transcript:
            raise ValueError("No speech could be transcribed.")

        return transcript


def transcribe_youtube_audio(video_id: str) -> str:
    """Download YouTube audio and transcribe it with Whisper."""

    ffmpeg_dir = (
        r"C:\Users\admin\Downloads\ffmpeg-9.0.2-essentials_build"
        r"\ffmpeg-9.0.2-essentials_build\bin"
    )

    video_url = f"https://www.youtube.com/watch?v={video_id}"

    with tempfile.TemporaryDirectory() as temp_dir:
        output_template = str(Path(temp_dir) / "audio.%(ext)s")

        ydl_options = {
            "format": "bestaudio/best",
            "outtmpl": output_template,
            "noplaylist": True,
            "quiet": True,
            "no_warnings": True,
            "ffmpeg_location": ffmpeg_dir,
            "postprocessors": [
                {
                    "key": "FFmpegExtractAudio",
                    "preferredcodec": "mp3",
                    "preferredquality": "128",
                }
            ],
        }

        with yt_dlp.YoutubeDL(ydl_options) as downloader:
            downloader.download([video_url])

        audio_path = Path(temp_dir) / "audio.mp3"

        if not audio_path.is_file():
            raise ValueError("Could not extract audio from the video.")

        model = WhisperModel(
            "base",
            device="cpu",
            compute_type="int8",
        )

        segments, _ = model.transcribe(
            str(audio_path),
            vad_filter=True,
        )

        transcript = "\n".join(
            segment.text.strip()
            for segment in segments
            if segment.text.strip()
        )

        if not transcript:
            raise ValueError("No speech could be transcribed.")

        return transcript

def transcribe_youtube_audio(video_id: str) -> str:
    """Download YouTube audio and transcribe it with Whisper."""

    ffmpeg_dir = (
        r"C:\Users\admin\Downloads\ffmpeg-9.0.2-essentials_build"
        r"\ffmpeg-9.0.2-essentials_build\bin"
    )

    video_url = f"https://www.youtube.com/watch?v={video_id}"

    with tempfile.TemporaryDirectory() as temp_dir:
        output_template = str(Path(temp_dir) / "audio.%(ext)s")

        ydl_options = {
            "format": "bestaudio/best",
            "outtmpl": output_template,
            "noplaylist": True,
            "quiet": True,
            "no_warnings": True,
            "ffmpeg_location": ffmpeg_dir,
            "postprocessors": [
                {
                    "key": "FFmpegExtractAudio",
                    "preferredcodec": "mp3",
                    "preferredquality": "128",
                }
            ],
        }

        with yt_dlp.YoutubeDL(ydl_options) as downloader:
            downloader.download([video_url])

        audio_path = Path(temp_dir) / "audio.mp3"

        if not audio_path.is_file():
            raise ValueError("Could not extract audio from the video.")

        model = WhisperModel(
            "base",
            device="cpu",
            compute_type="int8",
        )

        segments, _ = model.transcribe(
            str(audio_path),
            vad_filter=True,
        )

        transcript = "\n".join(
            segment.text.strip()
            for segment in segments
            if segment.text.strip()
        )

        if not transcript:
            raise ValueError("No speech could be transcribed.")

        return transcript

def transcribe_youtube_audio(video_id: str) -> str:
    """Download YouTube audio and transcribe it with Whisper."""

    ffmpeg_dir = (
        r"C:\Users\admin\Downloads\ffmpeg-9.0.2-essentials_build"
        r"\ffmpeg-9.0.2-essentials_build\bin"
    )

    video_url = f"https://www.youtube.com/watch?v={video_id}"

    with tempfile.TemporaryDirectory() as temp_dir:
        output_template = str(Path(temp_dir) / "audio.%(ext)s")

        ydl_options = {
            "format": "bestaudio/best",
            "outtmpl": output_template,
            "noplaylist": True,
            "quiet": True,
            "no_warnings": True,
            "ffmpeg_location": ffmpeg_dir,
            "postprocessors": [
                {
                    "key": "FFmpegExtractAudio",
                    "preferredcodec": "mp3",
                    "preferredquality": "128",
                }
            ],
        }

        with yt_dlp.YoutubeDL(ydl_options) as downloader:
            downloader.download([video_url])

        audio_path = Path(temp_dir) / "audio.mp3"

        if not audio_path.is_file():
            raise ValueError("Could not extract audio from the video.")

        model = WhisperModel(
            "base",
            device="cpu",
            compute_type="int8",
        )

        segments, _ = model.transcribe(
            str(audio_path),
            vad_filter=True,
        )

        transcript = "\n".join(
            segment.text.strip()
            for segment in segments
            if segment.text.strip()
        )

        if not transcript:
            raise ValueError("No speech could be transcribed.")

        return transcript



def transcribe_youtube_audio(video_id: str) -> str:
    """Download YouTube audio and transcribe it with Whisper."""

    ffmpeg_dir = (
        r"C:\Users\admin\Downloads\ffmpeg-9.0.2-essentials_build"
        r"\ffmpeg-9.0.2-essentials_build\bin"
    )

    video_url = f"https://www.youtube.com/watch?v={video_id}"

    with tempfile.TemporaryDirectory() as temp_dir:
        output_template = str(Path(temp_dir) / "audio.%(ext)s")

        ydl_options = {
            "format": "bestaudio/best",
            "outtmpl": output_template,
            "noplaylist": True,
            "quiet": True,
            "no_warnings": True,
            "ffmpeg_location": ffmpeg_dir,
            "postprocessors": [
                {
                    "key": "FFmpegExtractAudio",
                    "preferredcodec": "mp3",
                    "preferredquality": "128",
                }
            ],
        }

        with yt_dlp.YoutubeDL(ydl_options) as downloader:
            downloader.download([video_url])

        audio_path = Path(temp_dir) / "audio.mp3"

        if not audio_path.is_file():
            raise ValueError("Could not extract audio from the video.")

        model = WhisperModel(
            "base",
            device="cpu",
            compute_type="int8",
        )

        segments, _ = model.transcribe(
            str(audio_path),
            vad_filter=True,
        )

        transcript = "\n".join(
            segment.text.strip()
            for segment in segments
            if segment.text.strip()
        )

        if not transcript:
            raise ValueError("No speech could be transcribed.")

        return transcript

def transcribe_youtube_audio(video_id: str) -> str:
    """Download YouTube audio and transcribe it with timestamps."""

    ffmpeg_dir = (
        r"C:\Users\admin\Downloads\ffmpeg-9.0.2-essentials_build"
        r"\ffmpeg-9.0.2-essentials_build\bin"
    )

    video_url = f"https://www.youtube.com/watch?v={video_id}"

    with tempfile.TemporaryDirectory() as temp_dir:
        output_template = str(Path(temp_dir) / "audio.%(ext)s")

        ydl_options = {
            "format": "bestaudio/best",
            "outtmpl": output_template,
            "noplaylist": True,
            "quiet": True,
            "no_warnings": True,
            "ffmpeg_location": ffmpeg_dir,
            "postprocessors": [
                {
                    "key": "FFmpegExtractAudio",
                    "preferredcodec": "mp3",
                    "preferredquality": "128",
                }
            ],
        }

        with yt_dlp.YoutubeDL(ydl_options) as downloader:
            downloader.download([video_url])

        audio_path = Path(temp_dir) / "audio.mp3"

        if not audio_path.is_file():
            raise ValueError("Could not extract audio from the video.")

        model = WhisperModel(
            "base",
            device="cpu",
            compute_type="int8",
        )

        segments, _ = model.transcribe(
            str(audio_path),
            vad_filter=True,
        )

        lines = []

        for segment in segments:
            text = segment.text.strip()
            if not text:
                continue

            total_seconds = int(segment.start)
            minutes = total_seconds // 60
            seconds = total_seconds % 60
            timestamp = f"{minutes:02d}:{seconds:02d}"

            lines.append(f"[{timestamp}] {text}")

        if not lines:
            raise ValueError("No speech could be transcribed.")

        return "\n".join(lines)




def fetch_youtube_transcript(video_id: str) -> str:
    api = YouTubeTranscriptApi()

    try:
        transcript = api.fetch(
            video_id,
            languages=["hi", "en"]
        )
    except Exception:
        # If Hindi or English captions aren't available,
        # try another caption language offered by YouTube.
        available = api.list(video_id)
        first_transcript = next(iter(available), None)

        if first_transcript is None:
            raise ValueError("No captions are available for this video.")

        transcript = api.fetch(
            video_id,
            languages=[first_transcript.language_code]
        )

    lines = []

    for snippet in transcript:
        text = snippet.text.strip()
        if not text:
            continue

        total_seconds = int(snippet.start)
        minutes = total_seconds // 60
        seconds = total_seconds % 60
        timestamp = f"{minutes:02d}:{seconds:02d}"

        lines.append(f"[{timestamp}] {text}")

    if not lines:
        raise ValueError("The video transcript is empty.")

    return "\n".join(lines)


class WebpageTextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.text_parts = []
        self.skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "svg"}:
            self.skip_depth += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "svg"}:
            self.skip_depth = max(0, self.skip_depth - 1)

    def handle_data(self, data):
        if self.skip_depth == 0:
            text = data.strip()
            if text:
                self.text_parts.append(text)



def extract_pdf_visuals(file_path: Path) -> str:
    """Describe meaningful diagrams, charts, figures, and tables in a PDF."""
    print(f"Visual analysis called for: {file_path.name}")

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key or file_path.suffix.lower() != ".pdf":
        return ""

    descriptions = []

    try:
        client = genai.Client(api_key=api_key)

        with pymupdf.open(str(file_path)) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                print(f"Analyzing rendered PDF page {page_number}...")

                # Render the entire page, including vector diagrams and charts.
                pix = page.get_pixmap(
                    matrix=pymupdf.Matrix(1.5, 1.5),
                    alpha=False
                )

                image = Image.frombytes(
                    "RGB",
                    [pix.width, pix.height],
                    pix.samples
                )

                response = client.models.generate_content(
                    model=GEMINI_MODEL,
                    contents=[
                        (
                            "Analyze this rendered textbook or slide page. "
                            "Identify meaningful diagrams, charts, tables, "
                            "figures, and labeled illustrations. Describe their "
                            "labels, parts, relationships, and key takeaways "
                            "so a student can learn from them. Do not merely "
                            "summarize the surrounding text. Ignore decoration. "
                            "If there are no meaningful visuals, respond only "
                            "with NONE."
                        ),
                        image
                    ]
                )

                description = (response.text or "").strip()

                if description and description.upper() != "NONE":
                    descriptions.append(
                        f"[Visuals on PDF page {page_number}]\n{description}"
                    )
                    print(f"Visual description found on page {page_number}.")
                else:
                    print(f"No meaningful visuals found on page {page_number}.")

    except Exception as error:
        print(f"PDF visual analysis failed for {file_path.name}: {error}")

    return "\n\n".join(descriptions)



def extract_pptx_visuals(file_path: Path) -> str:
    """Describe embedded pictures in PowerPoint slides."""
    if file_path.suffix.lower() != ".pptx":
        return ""

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return ""

    descriptions = []

    try:
        client = genai.Client(api_key=api_key)
        presentation = Presentation(str(file_path))

        for slide_number, slide in enumerate(presentation.slides, start=1):
            for shape in slide.shapes:
                if shape.shape_type != 13:  # Embedded picture
                    continue

                print(
                    f"Analyzing picture on PowerPoint slide {slide_number}..."
                )

                image = Image.open(
                    io.BytesIO(shape.image.blob)
                ).convert("RGB")

                response = client.models.generate_content(
                    model=GEMINI_MODEL,
                    contents=[
                        (
                            "Describe this educational image for a student. "
                            "Explain meaningful diagrams, labels, charts, "
                            "figures, and relationships. Ignore decoration. "
                            "If it contains no meaningful educational visual, "
                            "respond only with NONE."
                        ),
                        image,
                    ],
                )

                description = (response.text or "").strip()

                if description and description.upper() != "NONE":
                    descriptions.append(
                        f"[Visual on PowerPoint slide {slide_number}]\n"
                        f"{description}"
                    )
                    print(
                        f"Visual description found on slide {slide_number}."
                    )

    except Exception as error:
        print(f"PowerPoint visual extraction failed: {error}")

    return "\n\n".join(descriptions)


def extract_docx_visuals(file_path: Path) -> str:
    """Describe embedded images in Word documents."""
    if file_path.suffix.lower() != ".docx":
        return ""

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return ""

    descriptions = []

    try:
        client = genai.Client(api_key=api_key)
        document = Document(str(file_path))

        image_parts = [
            part
            for part in document.part.related_parts.values()
            if part.content_type.startswith("image/")
        ]

        for image_number, image_part in enumerate(image_parts, start=1):
            print(f"Analyzing image {image_number} in Word document...")

            image = Image.open(
                io.BytesIO(image_part.blob)
            ).convert("RGB")

            response = client.models.generate_content(
                model=GEMINI_MODEL,
                contents=[
                    (
                        "Describe this educational image for a student. "
                        "Identify meaningful diagrams, charts, labels, "
                        "figures, and relationships. Explain what they show "
                        "and their key takeaways. Ignore decoration. "
                        "If there is no meaningful educational visual, "
                        "respond only with NONE."
                    ),
                    image,
                ],
            )

            description = (response.text or "").strip()

            if description and description.upper() != "NONE":
                descriptions.append(
                    f"[Visual {image_number} in Word document]\n"
                    f"{description}"
                )
                print(f"Visual description found for image {image_number}.")

    except Exception as error:
        print(f"Word visual extraction failed for {file_path.name}: {error}")

    return "\n\n".join(descriptions)



def extract_text(file_path: Path) -> str:
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        reader = PdfReader(str(file_path))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif extension == ".docx":
        document = Document(str(file_path))
        return "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
        )

    elif extension == ".pptx":
        presentation = Presentation(str(file_path))
        slide_texts = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text.strip()
            ]
            if texts:
                slide_texts.append(
                    f"Slide {slide_number}:\n" + "\n".join(texts)
                )
        return "\n\n".join(slide_texts)

    elif extension == ".txt":
        return file_path.read_text(encoding="utf-8", errors="replace")

    return ""
@app.get("/")
def home():
    return {
        "message": "Welcome to Knowvia!",
        "status": "Backend is running",
        "gemini_key_loaded": bool(os.getenv("GEMINI_API_KEY"))

    }

@app.get("/api/materials")
def get_materials():
    materials = []
    cache = load_material_cache()
    topic_store = load_topic_store()

    for file_path in UPLOAD_DIR.iterdir():
        if not file_path.is_file() or file_path.name.endswith(".part"):
            continue

        try:
            file_size = file_path.stat().st_size
        except OSError:
            continue

        saved = cache.get(file_path.name, {})

        # Text files such as saved webpage and YouTube transcripts
        # can be read directly without expensive document extraction.
        if file_path.suffix.lower() == ".txt" and not saved:
            try:
                text = file_path.read_text(encoding="utf-8")
                saved = {
                    "text_extracted": bool(text.strip()),
                    "text_preview": text[:300],
                    "text_length": len(text)
                }
                cache[file_path.name] = saved
            except (OSError, UnicodeError):
                saved = {}

        topic_data = topic_store.get(file_path.name, {})
        main_topics = topic_data.get("main_topics", [])

        materials.append({
            "filename": file_path.name,
            "size": file_size,
            "extension": file_path.suffix.lower(),
            "text_extracted": saved.get("text_extracted", False),
            "text_preview": saved.get("text_preview", ""),
            "extracted_topics": [
               topic.get("topic", "")
               for topic in main_topics
               if topic.get("topic")
            ]
       })

    save_material_cache(cache)
    return {"materials": materials}



 
@app.post("/api/upload")
async def upload_material(file: UploadFile = File(...)):
    filename = Path(file.filename or "").name
    extension = Path(filename).suffix.lower()

    if not filename or extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type."
        )

    destination = UPLOAD_DIR / filename
    temp_destination = UPLOAD_DIR / f"{filename}.part"
    size = 0

    try:
        with temp_destination.open("wb") as output:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)

                if size > MAX_FILE_SIZE:
                    raise HTTPException(
                        status_code=413,
                        detail="File exceeds the 200 MB limit."
                    )

                output.write(chunk)

        if size == 0:
            raise HTTPException(
                status_code=400,
                detail="The uploaded file is empty."
            )

        # Avoid saving duplicate copies of the same filename
        if destination.exists():
            temp_destination.unlink(missing_ok=True)

            cache = load_material_cache()
            existing_info = cache.get(filename, {})

            return {
                "message": "This file is already uploaded.",
                "filename": filename,
                "size": destination.stat().st_size,
                "status": "already_exists",
                "text_extracted": existing_info.get("text_extracted", False),
                "text_length": existing_info.get("text_length", 0),
                "text_preview": existing_info.get("text_preview", ""),
                "topic_analysis_status": existing_info.get(
                    "topic_analysis_status", "unknown"
                ),
            }

        temp_destination.replace(destination)

        

        extracted_text = extract_text(destination)
        visual_descriptions = extract_pdf_visuals(destination)
        visual_descriptions += "\n\n" + extract_pptx_visuals(destination)
        visual_descriptions += "\n\n" + extract_docx_visuals(destination)


        if visual_descriptions.strip():
            extracted_text += "\n\n" + visual_descriptions

        
        topic_analysis = analyze_material_topics(destination.name, extracted_text)
        
        cache = load_material_cache()
        topic_store = load_topic_store()
        cache[destination.name] = {
              "text_extracted": bool(extracted_text.strip()),
              "text_preview": extracted_text[:300],
              "text_length": len(extracted_text),
              "topic_analysis_status": "completed" if topic_analysis else "pending"
          }
        save_material_cache(cache)


        
        return {
    "message": "File uploaded successfully!",
    "filename": destination.name,
    "size": size,
    "status": "uploaded",
    "text_extracted": bool(extracted_text.strip()),
    "text_length": len(extracted_text),
    "text_preview": extracted_text[:300],
    "topic_analysis_status": "completed" if topic_analysis else "pending",
    }

    except Exception:
        if temp_destination.exists():
            temp_destination.unlink()
        raise

    finally:
        await file.close()


GEMINI_MODEL = "gemini-3.5-flash-lite"
TOPIC_STORE_PATH = Path(__file__).parent / "topic_store.json"
QUESTION_HISTORY_PATH = Path(__file__).parent / "question_history.json"


def load_topic_store():
    if not TOPIC_STORE_PATH.exists():
        return {}
    try:
        with TOPIC_STORE_PATH.open("r", encoding="utf-8") as file:
            data = json.load(file)
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def save_topic_store(data):
    with TOPIC_STORE_PATH.open("w", encoding="utf-8") as file:
        json.dump(data, file, indent=2, ensure_ascii=False)


def load_question_history():
    if not QUESTION_HISTORY_PATH.exists():
        return {}

    try:
        with open(QUESTION_HISTORY_PATH, "r", encoding="utf-8") as file:
            return json.load(file)
    except (json.JSONDecodeError, OSError):
        return {}


def save_question_history(history):
    with open(QUESTION_HISTORY_PATH, "w", encoding="utf-8") as file:
        json.dump(history, file, indent=2, ensure_ascii=False)


def normalize_question(question):
    return " ".join(question.lower().split())


def verify_quiz_answers(client, questions, study_material):
    verification_prompt = f"""
You are a strict answer verifier for an educational assessment.

Check every question against the provided study material.

For each question:
1. Verify that the correct answer is supported by the material.
2. For MCQs, verify that correctAnswer is the zero-based index
   of the correct option.
3. Check that the explanation agrees with the source.
4. Check that the cited page/location is plausible from the
   information provided. Do not invent missing evidence.
5. Reject questions that cannot be verified from the material.

Return ONLY a valid JSON array containing the verified questions.
Keep the original question fields unchanged for questions that pass.
Exclude questions that are incorrect, ambiguous, or unsupported.

Questions to verify:
{json.dumps(questions, ensure_ascii=False)}

Study material:
{study_material[:30000]}
"""

    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=verification_prompt
    )

    raw_text = (response.text or "").strip()

    if raw_text.startswith("```"):
        raw_text = re.sub(
            r"^```(?:json)?\s*|\s*```$",
            "",
            raw_text,
            flags=re.IGNORECASE
        ).strip()

    verified_questions = json.loads(raw_text)

    if not isinstance(verified_questions, list):
        raise ValueError("Answer verification did not return a JSON array.")

    return [
        question for question in verified_questions
        if isinstance(question, dict)
        and isinstance(question.get("question"), str)
    ]


def generate_assessment_report(client, results, study_material):
    prompt = f"""
You are Knowvia, an evidence-based AI learning companion.

Analyze the learner's assessment results using the study material.

Assessment results:
{json.dumps(results, ensure_ascii=False, indent=2)}

Study material:
{study_material[:20000]}

Return ONLY a valid JSON object with this structure:
{{
  "overallSummary": "Brief performance summary",
  "weakTopics": [
    {{
      "topic": "Topic name",
      "reason": "Evidence from the learner's incorrect answers",
      "misconception": "Likely misconception, or insufficient evidence",
      "recommendedRevision": "What the learner should revise",
      "source": "Relevant page or source location"
    }}
  ],
  "strengths": ["Topics answered correctly"],
  "nextSteps": ["Specific action the learner should take"]
}}

Rules:
- Base conclusions on the submitted results and study material.
- Focus weak topics on incorrect answers.
- Never claim a misconception with certainty from one answer.
- Do not invent page numbers or citations.
- If there is insufficient evidence, say so.
- Keep the report concise and actionable.
"""

    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=prompt
    )

    raw_text = (response.text or "").strip()

    if raw_text.startswith("```"):
        raw_text = re.sub(
            r"^```(?:json)?\s*|\s*```$",
            "",
            raw_text,
            flags=re.IGNORECASE
        ).strip()

    report = json.loads(raw_text)

    if not isinstance(report, dict):
        raise ValueError("Assessment report must be a JSON object.")

    return report




def analyze_material_topics(filename: str, text: str):
    if not text.strip():
        return None

    store = load_topic_store()

    # Reuse analysis if the saved material has not changed.
    if filename in store and store[filename].get("text_length") == len(text):
        return store[filename]

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return None
    
    
    question_history = load_question_history()
    previous_questions = question_history.get(filename, [])

    previous_questions_text = json.dumps(
    previous_questions,
    ensure_ascii=False,
    indent=2
)
    
    prompt = f"""
You are Knowvia's study-material analyzer.
Analyze the material below and return ONLY valid JSON in this structure:
{{
  "main_topics": [
    {{
      "topic": "Topic name",
      "subtopics": [
        {{
          "name": "Subtopic name",
          "key_concepts": ["Concept 1", "Concept 2"],
          "tags": ["tag1", "tag2"]
        }}
      ]
    }}
  ]
}}

Identify the important topics and subtopics actually covered.
Keep concepts concise. Do not invent information absent from the material.
Material filename: {filename}

Material:
{text[:18000]}
"""

    try:
        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt
        )

        result_text = response.text.strip()
        if result_text.startswith("```"):
            result_text = result_text.split("\n", 1)[-1]
            if result_text.endswith("```"):
                result_text = result_text[:-3].strip()

        analysis = json.loads(result_text)
        if not isinstance(analysis.get("main_topics"), list):
            return None

        record = {
            "text_length": len(text),
            "main_topics": analysis["main_topics"]
        }
        store[filename] = record
        save_topic_store(store)
        return record

    except Exception as error:
        print(f"Topic analysis failed for {filename}: {error}")
        return None


class TutorRequest(BaseModel):
    question: str
    subject: str = ""
    topic: str = ""
    language: str = "English"
    filename: str = ""


class QuizRequest(BaseModel):
    filename: str
    question_count: int = 5
    question_type: str = "mcq"
    difficulty: str = "Adaptive"
    language: str = "English"

    # Student-selected content scope
    selected_topics: list[str] = []
    selected_subtopics: list[str] = []

    # Assessment mode: "quiz" or "mock_exam"
    exam_mode: str = "quiz"

    # Optional question counts for a mixed assessment
    question_mix: dict[str, int] = {}


class URLMaterialRequest(BaseModel):
    url: str



def find_relevant_chunks(
    question: str,
    limit: int = 4,
    preferred_filename: str = ""
):
    stopwords = {
        "what", "are", "the", "with", "from", "that", "this",
        "does", "have", "your", "about", "into", "and", "for",
        "how", "can", "you", "my", "its"
    }

    def normalize(word):
        if len(word) > 4 and word.endswith("ies"):
            return word[:-3] + "y"
        if len(word) > 3 and word.endswith("s"):
            return word[:-1]
        return word

    words = {
        normalize(word)
        for word in re.findall(r"[a-z0-9]+", question.lower())
        if len(word) > 2 and word not in stopwords
    }

    if preferred_filename and (UPLOAD_DIR / preferred_filename).is_file():
        files_to_search = [UPLOAD_DIR / preferred_filename]
    else:
        files_to_search = list(UPLOAD_DIR.iterdir())

    scored = []

    for file_path in files_to_search:
        if not file_path.is_file() or file_path.name.endswith(".part"):
            continue

        try:
            text = extract_text(file_path)
        except Exception:
            continue

        if not text or not text.strip():
            continue

        filename_words = {
            normalize(word)
            for word in re.findall(r"[a-z0-9]+", file_path.stem.lower())
            if len(word) > 2
        }

        filename_score = len(words & filename_words) * 50

        for i in range(0, len(text), 1000):
            chunk = text[i:i + 1000]
            chunk_lower = chunk.lower()

            text_score = sum(
                1
                for word in words
                if re.search(rf"\b{re.escape(word)}\b", chunk_lower)
            )

            total_score = text_score + filename_score

            if preferred_filename and file_path.name == preferred_filename:
                total_score += 1000

            if total_score > 0:
                scored.append((total_score, file_path, chunk))

    scored.sort(key=lambda item: item[0], reverse=True)
    return scored[:limit]


class QuizSubmissionRequest(BaseModel):
    filename: str
    questions: list[dict]
    answers: dict[str, int | str]
    learner_id: str = "default_learner"


@app.post("/api/materials/url")
def add_url_material(request: URLMaterialRequest):
    url = request.url.strip()
    parsed = urlparse(url)

    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise HTTPException(
            status_code=400,
            detail="Please provide a valid HTTP or HTTPS URL."
        )
    
    video_id = get_youtube_video_id(url)
    if video_id:
        
        try:
            extracted_text = fetch_youtube_transcript(video_id)
        except Exception as caption_error:
            print("CAPTIONS ERROR:", repr(caption_error))

            try:
                extracted_text = transcribe_youtube_audio(video_id)
            except Exception as error:
                print("WHISPER ERROR:", repr(error))
                raise HTTPException(
                    status_code=422,
                    detail=(
                        "Could not transcribe this YouTube video. "
                        "Check the backend terminal for the actual errors."
                    )
                ) from error

        filename = f"youtube_{video_id}.txt"
        destination = UPLOAD_DIR / filename
        destination.write_text(extracted_text, encoding="utf-8")
        
        return {
            "message": "YouTube transcript added successfully!",
            "filename": filename,
            "text_length": len(extracted_text),
            "text_preview": extracted_text[:300]
        }


    import ipaddress
    import socket
    from urllib.error import HTTPError, URLError
    from urllib.request import HTTPRedirectHandler, build_opener

    class NoRedirectHandler(HTTPRedirectHandler):
        def redirect_request(self, req, fp, code, msg, headers, newurl):
            return None

    try:
        addresses = socket.getaddrinfo(
            parsed.hostname, None, type=socket.SOCK_STREAM
        )
        if not addresses:
            raise ValueError("Could not resolve website address.")

        for address in addresses:
            ip = ipaddress.ip_address(address[4][0])
            if not ip.is_global:
                raise ValueError("Private or local addresses are not allowed.")

        opener = build_opener(NoRedirectHandler())
        web_request = Request(
            url,
            headers={"User-Agent": "KnowviaStudyBot/1.0"}
        )

        with opener.open(web_request, timeout=10) as response:
            content_type = response.headers.get("Content-Type", "")
            if not any(
                kind in content_type.lower()
                for kind in ("text/html", "text/plain")
            ):
                raise HTTPException(
                    status_code=415,
                    detail="This link does not provide a readable webpage."
                )

            raw_content = response.read(2 * 1024 * 1024 + 1)
            if len(raw_content) > 2 * 1024 * 1024:
                raise HTTPException(
                    status_code=413,
                    detail="Webpage is too large to process."
                )

            charset = response.headers.get_content_charset() or "utf-8"
            page_content = raw_content.decode(charset, errors="replace")

        if "text/html" in content_type.lower():
            parser = WebpageTextParser()
            parser.feed(page_content)
            extracted_text = "\n".join(parser.text_parts)
        else:
            extracted_text = page_content

    except HTTPException:
        raise
    except (ValueError, OSError, URLError, HTTPError) as error:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read this webpage: {error}"
        )
    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail=f"Webpage processing failed: {error}"
        )

    if not extracted_text.strip():
        raise HTTPException(
            status_code=422,
            detail="No readable text was found on this webpage."
        )

    filename = f"webpage_{int(time.time())}.txt"
    destination = UPLOAD_DIR / filename
    counter = 1

    while destination.exists():
        filename = f"webpage_{int(time.time())}_{counter}.txt"
        destination = UPLOAD_DIR / filename
        counter += 1

    destination.write_text(extracted_text, encoding="utf-8")

    return {
        "message": "Webpage added successfully!",
        "filename": filename,
        "text_length": len(extracted_text),
        "text_preview": extracted_text[:300]
    }

@app.post("/api/tutor")
def ask_tutor(request: TutorRequest):
    question = request.question.strip()
    if not question:
        raise HTTPException(status_code=400, detail="Question is empty.")

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="Gemini API key is not configured.")

    

    if request.filename and (UPLOAD_DIR / request.filename).is_file():
        file_path = UPLOAD_DIR / request.filename
        full_text = extract_text(file_path)
        matches = [
            (1, file_path, full_text[i:i + 1000])
            for i in range(0, len(full_text), 1000)
        ]
    else:
        matches = find_relevant_chunks(
            question,
            preferred_filename=request.filename
        )

    if not matches:
        return {
            "content": "I couldn't find anything about this in your uploaded materials. "
                       "Try uploading notes or slides that cover this topic.",
            "citations": [],
            "covered": False,
        }


    excerpts = "\n\n".join(
        f"[Source: {path.name}]\n{chunk}" for _, path, chunk in matches
    )
    
    topic_store = load_topic_store()
    topic_context_parts = []

    for _, path, _ in matches:
        saved_analysis = topic_store.get(path.name)
        if saved_analysis:
            topic_context_parts.append(
                f"[Topic analysis for {path.name}]\n"
                + json.dumps(
                    saved_analysis.get("main_topics", []),
                    ensure_ascii=False
                )
            )

    topic_context = "\n\n".join(topic_context_parts)


    prompt = (
        "You are Knowvia, a friendly tutor for students. "
        "Prioritize the study material excerpts when they are relevant to the question. "
        "If the excerpts do not cover the question, explain the topic using your general knowledge. "
        "Clearly label such explanations as general knowledge, not information from the uploaded material. "
        "Explain simply, with a short example if it helps. "
        f"The required answer language is {request.language}. "
        "You MUST write the entire answer in this language, regardless of the language used in the video or source material. "
        "Never switch to Hindi or Hinglish unless the required answer language is Hindi or Hinglish. "
        "If the requested language is Hinglish, use natural Hindi-English mixed speech in Roman script. "
        "If the excerpts do not contain the answer, say so clearly.\n\n"
        f"Subject: {request.subject}\nTopic: {request.topic}\n\n"
        f"Study material:\n{excerpts}\n\n"
        f"Saved topic analysis (use as organization hints only):\n{topic_context}\n\n"
        f"Question: {question}"
    )

    try:
        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(model=GEMINI_MODEL, contents=prompt)
    except Exception as error:
        raise HTTPException(status_code=502, detail=f"Gemini request failed: {error}")

    
    sources = []
    for _, path, _ in matches:
        if path.name not in sources:
            sources.append(path.name)

    
    return {
        "content": response.text,
        "citations": [
            {
                "source": name,
                "type": "PDF" if name.lower().endswith(".pdf") else "Document"
                }
                  for name in sources
                  ],
                  "covered": True,
                  "contexts": [chunk for _, _, chunk in matches],
                  }





@app.post("/api/quiz/generate")
def generate_quiz(request: QuizRequest):
    if not request.filename.strip():
        raise HTTPException(
            status_code=400,
            detail="Please select a study material first."
        )

    file_path = UPLOAD_DIR / request.filename

    if not file_path.is_file():
        raise HTTPException(
            status_code=404,
            detail="The selected study material was not found."
        )

    try:
        full_text = extract_text(file_path)
    except Exception as error:
        raise HTTPException(
            status_code=422,
            detail=f"Could not read the study material: {error}"
        )

    if not full_text or not full_text.strip():
        raise HTTPException(
            status_code=422,
            detail="No readable text was found in this material."
        )
    
    # Load the saved topic and subtopic structure.
    topic_analysis = analyze_material_topics(
        file_path.name,
        full_text
    )

    main_topics = (
        topic_analysis.get("main_topics", [])
        if topic_analysis
        else []
    )

    selected_scope = []

    for topic in main_topics:
        topic_name = topic.get("topic", "")
        subtopics = topic.get("subtopics", [])

        # Include a whole topic if it was selected.
        if (
            not request.selected_topics
            and not request.selected_subtopics
        ) or topic_name in request.selected_topics:
            selected_scope.append({
                "topic": topic_name,
                "subtopics": subtopics
            })
            continue

        # Otherwise, include only the selected subtopics.
        matching_subtopics = [
            subtopic
            for subtopic in subtopics
            if subtopic.get("name", "")
            in request.selected_subtopics
        ]

        if matching_subtopics:
            selected_scope.append({
                "topic": topic_name,
                "subtopics": matching_subtopics
            })

    # Reject selections that don't match the saved topic analysis.
    if (
        request.selected_topics or request.selected_subtopics
    ) and not selected_scope:
        raise HTTPException(
            status_code=400,
            detail=(
                "The selected topics were not found in this material. "
                "Please refresh the topic list and try again."
            )
        )


    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(
            status_code=500,
            detail="Gemini API key is not configured."
        )
    
    scope_context = json.dumps(
        selected_scope,
        ensure_ascii=False,
        indent=2
    )


    
   
    question_type = request.question_type.lower()
    exam_mode = request.exam_mode.lower()
    question_mix = request.question_mix or {}

    if exam_mode not in {"quiz", "mock_exam"}:
        raise HTTPException(
            status_code=400,
            detail="Assessment mode must be 'quiz' or 'mock_exam'."
        )

    if exam_mode == "mock_exam":
        question_count = sum(question_mix.values())
    else:
        question_count = max(1, min(request.question_count, 10))

    question_mix = request.question_mix or {}

    if exam_mode == "mock_exam":
        allowed_types = {"mcq", "short", "numerical"}

        if any(
            kind not in allowed_types
            or not isinstance(count, int)
            or isinstance(count, bool)
            or count < 0
            for kind, count in question_mix.items()
        ):
            raise HTTPException(
                status_code=400,
                detail="Invalid question mix."
            )

        if not question_mix or sum(question_mix.values()) == 0:
            raise HTTPException(
                status_code=400,
                detail=(
                    "For a mock exam, provide counts for at least "
                    "one question type."
                )
            )

        if sum(question_mix.values()) > 10:
            raise HTTPException(
                status_code=400,
                detail="A mock exam can contain at most 10 questions."
            )

    difficulty = request.difficulty

    if question_type not in {"mcq", "short", "numerical"}:
        raise HTTPException(
            status_code=400,
            detail="Unsupported question type."
        )
    question_history = load_question_history()
    previous_questions = question_history.get(file_path.name, [])
    previous_questions_text = json.dumps(
        previous_questions,
        ensure_ascii=False,
        indent=2
     )

    prompt = f"""
    
{"Mock exam question counts: " + json.dumps(question_mix)
 if exam_mode == "mock_exam"
 else "This is a regular quiz with one question type."}

For a mock exam, generate exactly the requested number of
questions for each type in the question mix. Include every
requested type and do not replace one type with another.

For a regular quiz, generate questions only of the requested
question type.

You are Knowvia, an educational quiz generator.

Create exactly {question_count} original quiz questions using
ONLY the study material provided below.



Question type: {question_type}
Difficulty: {difficulty}
Language: {request.language}
Assessment mode: {request.exam_mode}

Generate questions ONLY from the following student-selected scope.
If the scope is empty, use the available study material topics.

Selected topics and subtopics:
{scope_context}

Previously used questions from earlier assessments:
{previous_questions_text}

IMPORTANT: Do not repeat any previously used question.
Generate new questions testing different aspects of the selected topics.
Avoid paraphrasing a previous question to make it appear new.
If there are too few unused concepts, prioritize genuinely new questions
and do not invent facts outside the study material.



Return ONLY a valid JSON array. Do not include Markdown fences
or any text outside the JSON.


Every question must have these fields:
- id: unique string
- question: question text
- type: the applicable question type
- topic: the main topic this question tests
- subtopic: the specific subtopic, or an empty string if unavailable
- difficulty: "Easy", "Medium", or "Hard"
- explanation: a brief explanation grounded in the material
- source: object containing:
  - source: exact filename
  - type: "PDF", "PowerPoint", "Word", or "Document"
  - location: exact page number, slide number, or "Unavailable"
  
Only provide source locations that can be verified from the material.
Never invent page or slide numbers. If the exact location is unavailable,
use "Unavailable".


For MCQs, also include:
- options: array of exactly four strings
- correctAnswer: zero-based index of the correct option

For short-answer questions, also include:
- acceptableAnswers: array of one or more acceptable answer strings

For numerical questions, also include:
- numericAnswer: correct numerical answer
- numericTolerance: acceptable numerical difference, normally 0

If difficulty is "Adaptive", generate a balanced mixture of
Easy, Medium, and Hard questions. Otherwise, use the requested
difficulty for every question.

Do not invent facts that are absent from the study material.
If the material is insufficient for the requested number of
good questions, return only the questions it supports.

Filename: {file_path.name}

Study material:
{full_text[:30000]}
"""

    try:
        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt
        )

        raw_text = (response.text or "").strip()
        if raw_text.startswith("```"):
            raw_text = re.sub(
                r"^```(?:json)?\s*|\s*```$",
                "",
                raw_text,
                flags=re.IGNORECASE
            ).strip()

        questions = json.loads(raw_text)

        if not isinstance(questions, list):
            raise ValueError("The response was not a JSON array.")
        
        allowed_types = (
            set(question_mix.keys())
            if exam_mode == "mock_exam"
            else {question_type}
        )
        questions = [
            question for question in questions
            if isinstance(question, dict)
            and isinstance(question.get("question"), str)
            and question.get("type") in allowed_types
        ]

        # Remove questions already used in previous assessments
        used_questions = {
            normalize_question(item)
            for item in previous_questions
            if isinstance(item, str)
        }

        unique_questions = []

        for question in questions:
            question_text = question.get("question", "").strip()
            normalized = normalize_question(question_text)

            if not normalized or normalized in used_questions:
                continue

            used_questions.add(normalized)
            unique_questions.append(question)

        questions = unique_questions
        
        
        source_type = {
            ".pdf": "PDF",
            ".pptx": "PowerPoint",
            ".docx": "Word",
        }.get(file_path.suffix.lower(), "Document")

        for index, question in enumerate(questions):
            question["id"] = str(question.get("id") or index + 1)

            question["topic"] = str(
                question.get("topic") or "Unspecified"
            )
            question["subtopic"] = str(
                question.get("subtopic") or ""
            )

            if question.get("difficulty") not in {
                "Easy", "Medium", "Hard"
            }:
                question["difficulty"] = (
                    difficulty
                    if difficulty in {"Easy", "Medium", "Hard"}
                    else "Medium"
                )

            generated_source = question.get("source")
            location = (
                generated_source.get("location", "Unavailable")
                if isinstance(generated_source, dict)
                else "Unavailable"
            )

            question["source"] = {
                "source": file_path.name,
                "type": source_type,
                "location": location
            }

       
        questions = verify_quiz_answers(
            client,
            questions,
            full_text
)
        if not questions:
            raise ValueError(
                "No new, unused questions were generated. "
                "Try a different topic or upload new material."
            )
        
        # Save newly generated questions to history
        for question in questions:
            question_text = question.get("question", "").strip()

            if question_text:
                previous_questions.append(question_text)

        question_history[file_path.name] = previous_questions
        save_question_history(question_history)

        return {
            "filename": file_path.name,
            "questions": questions
        }

    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(
            status_code=502,
            detail=f"Quiz generation failed: {error}"
        )
    
@app.get("/api/topics")
def get_topics():
    return load_topic_store()


@app.post("/api/quiz/submit")
def submit_quiz(request: QuizSubmissionRequest):
    
    file_path = UPLOAD_DIR / request.filename

    if not file_path.is_file():
        raise HTTPException(
            status_code=404,
            detail="The selected study material was not found."
        )

    try:
        full_text = extract_text(file_path)
    except Exception as error:
        raise HTTPException(
            status_code=422,
            detail=f"Could not read the study material: {error}"
        )

    if not full_text or not full_text.strip():
        raise HTTPException(
            status_code=422,
            detail="No readable text was found in this material."
        )

    results = []
    correct_count = 0
    learner_store = load_learner_store()
    learner = get_or_create_learner(
    learner_store,
    request.learner_id.strip() or "default_learner"
)


    for question in request.questions:
        question_id = question.get("id")
        user_answer = request.answers.get(question_id)
        correct_answer = question.get("correctAnswer")

        is_correct = (
            user_answer == correct_answer
            or (
                isinstance(user_answer, str)
                and isinstance(correct_answer, int)
                and user_answer.strip().isdigit()
                and int(user_answer.strip()) == correct_answer
            )
        )

        if is_correct:
            correct_count += 1
            topic = question.get("topic", "General")
            update_topic_mastery(
                learner,
                topic,
                is_correct,
                "quiz"
            )


        results.append({
            "questionId": question_id,
            "question": question.get("question"),
            "topic": question.get("topic", "General"),
            "subtopic": question.get("subtopic", "General"),
            "userAnswer": user_answer,
            "correctAnswer": correct_answer,
            "isCorrect": is_correct,
            "feedback": (
                (
                    "Correct! " if is_correct else "Not quite. "
                )
                + question.get("explanation", "")
                + (
                    f" Source: {question.get('source', {}).get('source', 'Study material')}, "
                    f"{question.get('source', {}).get('location', 'location unavailable')}."
                )
            ),

            "explanation": question.get("explanation", ""),
            "source": question.get("source", {})
        })

    
    total = len(results)
    learner["quizHistory"].append({
        "filename": request.filename,
        "score": correct_count,
        "totalQuestions": total,
        "percentage": round(correct_count / total * 100, 2) if total else 0,
        "timestamp": time.strftime("%Y-%m-%dT%H:%M:%S")
    })
    save_learner_store(learner_store)

    try:
        api_key = os.getenv("GEMINI_API_KEY")
        if not api_key:
            raise HTTPException(
            status_code=500,
            detail="Gemini API key is not configured."
    )

        client = genai.Client(api_key=api_key)

        report = generate_assessment_report(
            client,
            results,
            full_text  
    )
    except Exception as error:
        raise HTTPException(
            status_code=502,
            detail=f"Could not generate assessment report: {error}"
        )

    return {
        "filename": request.filename,
        "score": correct_count,
        "totalQuestions": total,
        "percentage": round(correct_count / total * 100, 2) if total else 0,
        "results": results,
        "report": report
    }

@app.get("/api/learners/{learner_id}")
def get_learner_profile(learner_id: str):
    if not learner_id.strip():
        raise HTTPException(
            status_code=400,
            detail="Learner ID cannot be empty."
        )

    learner_store = load_learner_store()

    learner = learner_store.get(learner_id)

    if learner is None:
        return {
            "learnerId": learner_id,
            "isNewLearner": True,
            "topicMastery": {},
            "quizHistory": [],
            "conversationHistory": [],
            "message": "No learning history yet. A diagnostic quiz is recommended."
        }

    return learner



